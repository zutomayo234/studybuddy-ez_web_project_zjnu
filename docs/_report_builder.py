"""
生成五份实验报告（.docx）
严格套用《实验报告参考模板00931.doc》结构：
  封面 → 目录 → 实验N：1.项目介绍 2.相关开发技术 3.项目设计 4.关键代码 5.运行结果 6.心得（使用AI情况）
运行结果小节插入"截图占位"，由本人替换为实际运行截图。
"""

import os
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

OUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "实验报告")
os.makedirs(OUT_DIR, exist_ok=True)
STUDENT_NO = "20243190415"
STUDENT_NAME = "方野"
CLASS_NAME = "（请填写班级）"
COURSE_NO = "00931"

CN_FONT = "宋体"
HEAD_FONT = "黑体"
CODE_FONT = "Consolas"


# ----------------------------------------------------------------------
# 基础样式工具
# ----------------------------------------------------------------------
def set_run_font(run, cn=CN_FONT, en="Times New Roman", size=12, bold=False, color=None):
    run.font.name = en
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), en)
    rfonts.set(qn("w:hAnsi"), en)
    rfonts.set(qn("w:eastAsia"), cn)


def shade(paragraph, fill="F2F2F2"):
    """给段落加底纹（用于代码块）"""
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def add_para(doc, text="", size=12, bold=False, cn=CN_FONT, align=None,
             space_before=0, space_after=6, indent_chars=0, line_spacing=1.5):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    if indent_chars:
        pf.first_line_indent = Pt(size * indent_chars)
    if text:
        run = p.add_run(text)
        set_run_font(run, cn=cn, size=size, bold=bold)
    return p


def add_heading(doc, text, level=1):
    sizes = {1: 16, 2: 14, 3: 13}
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(14 if level == 1 else 10)
    pf.space_after = Pt(8)
    run = p.add_run(text)
    set_run_font(run, cn=HEAD_FONT, en="Arial", size=sizes.get(level, 12), bold=True)
    return p


def add_body(doc, text, indent=True):
    """正文段落，首行缩进 2 字符"""
    return add_para(doc, text, size=12, indent_chars=2 if indent else 0, space_after=6)


def add_bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.4
        run = p.add_run(it)
        set_run_font(run, size=12)


def add_code(doc, code, caption=None):
    """插入代码块（等宽字体 + 灰色底纹 + 单倍行距）"""
    if caption:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(caption)
        set_run_font(run, size=10.5, bold=True, color=(0x40, 0x40, 0x40))
    for line in code.rstrip("\n").split("\n"):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        pf.left_indent = Pt(6)
        run = p.add_run(line if line else " ")
        set_run_font(run, cn=CN_FONT, en=CODE_FONT, size=9)
        shade(p)
    # 代码块后留白
    add_para(doc, "", size=6, space_after=2)


def add_table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        set_run_font(run, cn=HEAD_FONT, size=10.5, bold=True)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            run = p.add_run(str(v))
            set_run_font(run, size=10.5)
    if widths:
        for r in t.rows:
            for i, w in enumerate(widths):
                r.cells[i].width = Cm(w)
    add_para(doc, "", size=6, space_after=2)
    return t


def add_screenshot_placeholder(doc, idx, what):
    """运行结果截图占位"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(6)
    pf.space_after = Pt(4)
    pf.line_spacing = 1.3
    run = p.add_run(f"【截图 {idx} 占位】")
    set_run_font(run, cn=HEAD_FONT, size=11, bold=True, color=(0xC0, 0x00, 0x00))
    run2 = p.add_run(what)
    set_run_font(run2, size=11, color=(0xC0, 0x00, 0x00))

    # 占位方框
    box = doc.add_table(rows=1, cols=1)
    box.style = "Table Grid"
    cell = box.rows[0].cells[0]
    cell.text = ""
    bp = cell.paragraphs[0]
    bp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    br = bp.add_run("\n\n\n（此处粘贴运行截图）\n\n\n")
    set_run_font(br, size=11, color=(0x99, 0x99, 0x99))
    add_para(doc, "", size=6, space_after=4)


# ----------------------------------------------------------------------
# 封面 + 目录
# ----------------------------------------------------------------------
def build_cover(doc, exp_no, exp_title):
    for _ in range(2):
        add_para(doc, "", size=12, space_after=0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("浙江师范大学")
    set_run_font(run, cn=HEAD_FONT, size=18, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("《团队规范项目实训》")
    set_run_font(run, cn=HEAD_FONT, size=22, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("个人项目实验报告")
    set_run_font(run, cn=HEAD_FONT, size=22, bold=True)

    add_para(doc, "", size=12, space_after=0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"实验{exp_no}　{exp_title}")
    set_run_font(run, cn=HEAD_FONT, size=16, bold=True, color=(0x1F, 0x3B, 0x73))

    for _ in range(6):
        add_para(doc, "", size=12, space_after=0)

    for label, value in [("学　　号：", STUDENT_NO), ("姓　　名：", STUDENT_NAME), ("班　　级：", CLASS_NAME)]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"{label}{value}")
        set_run_font(run, size=14)

    add_para(doc, "", size=12, space_after=0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(f"课程编号：{COURSE_NO}　　")
    set_run_font(run, size=14)

    doc.add_page_break()


def build_toc(doc, exp_no, exp_title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("目　　录")
    set_run_font(run, cn=HEAD_FONT, size=18, bold=True)
    add_para(doc, "", size=12, space_after=6)

    p = doc.add_paragraph()
    run = p.add_run(f"实验{exp_no}　{exp_title}")
    set_run_font(run, cn=HEAD_FONT, size=13, bold=True)
    for item in ["1. 项目介绍", "2. 相关开发技术", "3. 项目设计（概要设计 + 详细设计 + 数据库设计）",
                 "4. 关键代码", "5. 运行结果", "6. 心得（使用 AI 情况）"]:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(24)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(item)
        set_run_font(run, size=12)

    doc.add_page_break()


# ----------------------------------------------------------------------
# 合并报告（五份实验报告合成一份，套用参考模板结构）
# ----------------------------------------------------------------------
TOC_SECTIONS = [
    "1. 项目介绍",
    "2. 相关开发技术",
    "3. 项目设计（包括概要设计 + 详细设计 + 数据库设计）",
    "4. 关键代码",
    "5. 运行结果",
    "6. 心得（使用 AI 情况）",
]


def build_merged_cover(doc, subtitle="（实验一 ～ 实验五 · 学习搭子 StudyBuddy）"):
    for _ in range(2):
        add_para(doc, "", size=12, space_after=0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("浙江师范大学")
    set_run_font(run, cn=HEAD_FONT, size=18, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("《团队规范项目实训》")
    set_run_font(run, cn=HEAD_FONT, size=22, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("个人项目实验报告")
    set_run_font(run, cn=HEAD_FONT, size=22, bold=True)

    add_para(doc, "", size=12, space_after=0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(subtitle)
    set_run_font(run, cn=HEAD_FONT, size=14, bold=True, color=(0x1F, 0x3B, 0x73))

    for _ in range(6):
        add_para(doc, "", size=12, space_after=0)

    for label, value in [("学　　号：", STUDENT_NO), ("姓　　名：", STUDENT_NAME), ("班　　级：", CLASS_NAME)]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"{label}{value}")
        set_run_font(run, size=14)

    add_para(doc, "", size=12, space_after=0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(f"课程编号：{COURSE_NO}　　")
    set_run_font(run, size=14)

    doc.add_page_break()


def build_merged_toc(doc, reports):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("目　　录")
    set_run_font(run, cn=HEAD_FONT, size=18, bold=True)
    add_para(doc, "", size=12, space_after=6)

    for r in reports:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(6)
        pf.space_after = Pt(3)
        run = p.add_run(f"实验{r['no']}　{r['title']}")
        set_run_font(run, cn=HEAD_FONT, size=13, bold=True)
        for s in TOC_SECTIONS:
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Pt(24)
            pf.space_after = Pt(2)
            run = p.add_run(s)
            set_run_font(run, size=11.5)

    doc.add_page_break()


def finish_merged(doc, reports, filename):
    build_merged_cover(doc)
    build_merged_toc(doc, reports)

    for i, r in enumerate(reports):
        if i > 0:
            doc.add_page_break()

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"实验{r['no']}　{r['title']}")
        set_run_font(run, cn=HEAD_FONT, size=17, bold=True)
        add_para(doc, "", size=8, space_after=4)

        add_heading(doc, "1. 项目介绍", 1)
        for para in r['intro']:
            add_body(doc, para)

        add_heading(doc, "2. 相关开发技术", 1)
        r['tech'](doc)

        add_heading(doc, "3. 项目设计（包括概要设计 + 详细设计 + 数据库设计）", 1)
        r['design'](doc)

        add_heading(doc, "4. 关键代码", 1)
        r['code'](doc)

        add_heading(doc, "5. 运行结果", 1)
        r['run'](doc)

        add_heading(doc, "6. 心得（使用 AI 情况）", 1)
        r['reflect'](doc)

    path = filename if os.path.isabs(filename) else os.path.join(OUT_DIR, filename)
    doc.save(path)
    print(f"已生成：{path}")
    return path


def new_doc():
    doc = Document()
    # 页边距
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.17)
        section.right_margin = Cm(3.17)
    # 正文默认样式
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
    return doc


def finish_recent(doc, exp_no, exp_title, filename, common_intro, tech, design, code, run_result, reflection):
    """统一的报告装配流程"""
    build_cover(doc, exp_no, exp_title)
    build_toc(doc, exp_no, exp_title)

    # 实验标题
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"实验{exp_no}　{exp_title}")
    set_run_font(run, cn=HEAD_FONT, size=17, bold=True)
    add_para(doc, "", size=8, space_after=4)

    add_heading(doc, "1. 项目介绍", 1)
    for para in common_intro:
        add_body(doc, para)

    add_heading(doc, "2. 相关开发技术", 1)
    tech(doc)

    add_heading(doc, "3. 项目设计（包括概要设计 + 详细设计 + 数据库设计）", 1)
    design(doc)

    add_heading(doc, "4. 关键代码", 1)
    code(doc)

    add_heading(doc, "5. 运行结果", 1)
    run_result(doc)

    add_heading(doc, "6. 心得（使用 AI 情况）", 1)
    reflection(doc)

    path = os.path.join(OUT_DIR, filename)
    doc.save(path)
    print(f"已生成：{path}")
    return path
